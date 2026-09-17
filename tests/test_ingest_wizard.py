"""Wizard decisions are Enter-to-accept; --one-shot writes and stops."""
from pathlib import Path

from wing_parser.classifier.provider import ProviderError
from wing_parser.cli.__main__ import main
from wing_parser.showcontext.ingest.wizard import run_wizard

DATA = Path(__file__).resolve().parent / "data"
VIVO = str(DATA / "05102022_vivo_ Event Rundown.xlsx")
MAP_NAME = "05102022_vivo_ Event Rundown.map.yaml"


class VivoProvider:
    """The proposal Task 9's suite already proved clean end-to-end."""

    def complete_json(self, system, user, schema):
        return {
            "sheet": "Rundown",
            "header_row": 4,
            "columns": '{"id": "B", "time": "C", "title": "F"}',
            "headers": '{"performers": "On stage", "note": "CHUẨN BỊ"}',
        }


class DeadProvider:
    def complete_json(self, system, user, schema):
        raise ProviderError("key missing")


class BothProvider:
    """Mapping proposals AND term guesses from one offline fake.

    The term-guess system prompt names kinds; the mapping-proposal one
    asks for sheet/columns/headers -- the schema tells them apart.
    """

    def complete_json(self, system, user, schema):
        if "sheet" in schema.get("properties", {}):
            return {
                "sheet": "Rundown",
                "header_row": 1,
                "columns": '{"id": "A", "time": "", "title": "B"}',
                "headers": '{"performers": "On stage"}',
            }
        return {"kind": "speech.playback", "confidence": 0.8}


def _fake_provider(monkeypatch, provider):
    """The wizard builds its own provider offline; swap the factory."""
    monkeypatch.setattr(
        "wing_parser.classifier.provider.make_provider", lambda config: provider
    )
    monkeypatch.delenv("WING_PROVIDER_CONFIG", raising=False)


def test_all_enter_accepts_proposal_and_imports(tmp_path, capsys, monkeypatch):
    _fake_provider(monkeypatch, VivoProvider())
    monkeypatch.chdir(tmp_path)
    out = tmp_path / "vivo.yaml"
    code = run_wizard(
        VIVO,
        input_fn=lambda *a, **k: "",
        print_fn=lambda *a, **k: None,
        output=str(out),
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 0
    assert out.exists()
    assert "segments:" in out.read_text(encoding="utf-8")
    assert (tmp_path / MAP_NAME).exists()

    from wing_parser.showcontext import load_show_context

    context = load_show_context(out)
    assert len(context.segments) > 0
    # The point of the proposal: performers resolved offline, not every
    # segment left as an unreadable comment.
    assert any(segment.expects for segment in context.segments)


def test_a_missing_workbook_is_one_error_line_not_a_traceback(
    tmp_path, capsys, monkeypatch
):
    _fake_provider(monkeypatch, VivoProvider())
    monkeypatch.chdir(tmp_path)
    missing = str(tmp_path / "nope.xlsx")
    asked = []
    code = run_wizard(
        missing,
        input_fn=lambda *a, **k: (asked.append(a), "")[1],
        print_fn=lambda *a, **k: print(*a),
        output=None,
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 1
    out = capsys.readouterr().out
    assert "Traceback" not in out
    assert "nope.xlsx" in out
    assert not asked  # refused before a single question was asked


def test_a_corrupt_workbook_degrades_to_the_manual_fallback_line(
    tmp_path, capsys, monkeypatch
):
    """An existing but non-xlsx file dies inside sampling, before the
    provider is ever called -- still one line, never a traceback."""
    _fake_provider(monkeypatch, VivoProvider())
    monkeypatch.chdir(tmp_path)
    fake = tmp_path / "broken.xlsx"
    fake.write_text("this is not a real spreadsheet\n", encoding="utf-8")
    answers = iter(["Rundown", "4", "B", "C", "F", "", "", "", "", ""])
    code = run_wizard(
        str(fake),
        input_fn=lambda *a, **k: next(answers),
        print_fn=lambda *a, **k: print(*a),
        output=str(tmp_path / "out.yaml"),
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 1
    captured = capsys.readouterr()
    assert "Traceback" not in captured.out and "Traceback" not in captured.err
    lines = [
        line
        for line in captured.out.splitlines()
        if "no model assist" in line
    ]
    assert len(lines) == 1


def test_one_shot_writes_only_the_map(tmp_path, capsys, monkeypatch):
    _fake_provider(monkeypatch, VivoProvider())
    monkeypatch.chdir(tmp_path)
    out = tmp_path / "vivo.yaml"
    code = run_wizard(
        VIVO,
        input_fn=lambda *a, **k: "",
        print_fn=lambda *a, **k: print(*a),
        output=str(out),
        force=False,
        scene=None,
        one_shot=True,
    )
    assert code == 0
    maps = list(tmp_path.glob("*.map.yaml"))
    assert len(maps) == 1
    assert maps[0].name == MAP_NAME
    assert not out.exists()
    assert ".map.yaml" in capsys.readouterr().out


def test_closed_input_is_one_line_and_a_clean_nonzero_exit(tmp_path, capsys, monkeypatch):
    """An exhausted stdin (Ctrl+Z on Windows) must not traceback through
    _ask -- the module promises never a traceback. An empty iterator is
    the scripted shape of that: every question hits a dead input_fn."""
    _fake_provider(monkeypatch, VivoProvider())
    monkeypatch.chdir(tmp_path)
    answers = iter([])
    code = run_wizard(
        VIVO,
        input_fn=lambda *a, **k: next(answers),
        print_fn=lambda *a, **k: print(*a),
        output=str(tmp_path / "out.yaml"),
        force=False,
        scene=None,
        one_shot=False,
    )
    captured = capsys.readouterr()
    assert code == 1
    assert "Traceback" not in captured.out and "Traceback" not in captured.err
    assert "input closed" in captured.out
    assert not list(tmp_path.glob("*.map.yaml"))  # nothing was saved


def test_kill_switch_skips_the_mapping_proposal(tmp_path, capsys, monkeypatch):
    """WING_DISABLE_LLM=1 must gate the wizard's model paths too, not only
    classifier/llm.py: no provider is ever constructed, one line, manual
    questions with no defaults."""
    monkeypatch.setenv("WING_DISABLE_LLM", "1")

    def forbidden(config):
        raise AssertionError("provider must not be constructed")

    monkeypatch.setattr(
        "wing_parser.classifier.provider.make_provider", forbidden
    )
    monkeypatch.delenv("WING_PROVIDER_CONFIG", raising=False)
    monkeypatch.chdir(tmp_path)
    answers = iter([
        "Rundown", "4", "B", "C", "F", "On stage", "CHUẨN BỊ", "", "", "",
    ])
    out = tmp_path / "vivo.yaml"
    code = run_wizard(
        VIVO,
        input_fn=lambda *a, **k: next(answers),
        print_fn=lambda *a, **k: print(*a),
        output=str(out),
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 0
    lines = [
        line for line in capsys.readouterr().out.splitlines()
        if "WING_DISABLE_LLM" in line
    ]
    # One line per gated call site: the proposal and the term guess
    # (the vivo sheet has unresolved performers, so both paths run).
    assert any("continuing manually" in line for line in lines)
    assert any("stay as comments" in line for line in lines)
    assert out.exists()


def test_kill_switch_skips_the_term_guess(tmp_path, capsys, monkeypatch):
    """With the switch on, unresolved terms stay comments: no provider is
    built and no Record question is ever asked."""
    from openpyxl import Workbook

    monkeypatch.setenv("WING_DISABLE_LLM", "1")

    def forbidden(config):
        raise AssertionError("provider must not be constructed")

    monkeypatch.setattr(
        "wing_parser.classifier.provider.make_provider", forbidden
    )
    monkeypatch.delenv("WING_PROVIDER_CONFIG", raising=False)
    monkeypatch.chdir(tmp_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rundown"
    ws.append(["No", "Nội dung"])
    ws.append([1, "Đón khách", "tốp múa"])
    xlsx = tmp_path / "run.xlsx"
    wb.save(xlsx)

    asked = []
    answers = iter(["Rundown", "1", "A", "", "B", "", "", "", "", ""])

    def answer(prompt):
        asked.append(prompt)
        return next(answers)

    out = tmp_path / "out.yaml"
    code = run_wizard(
        str(xlsx),
        input_fn=answer,
        print_fn=lambda *a, **k: print(*a),
        output=str(out),
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 0
    captured = capsys.readouterr()
    assert not any(p.startswith("Record") for p in asked)
    assert "WING_DISABLE_LLM" in captured.out


class UnverifiedProvider:
    """A proposal that survives decoding but fails the title check.

    propose_mapping retries once and still gets problems back, so the
    wizard receives a MappingProposal with a non-empty problems tuple.
    """

    def complete_json(self, system, user, schema):
        return {
            "sheet": "Rundown",
            "header_row": 4,
            "columns": '{"id": "B", "time": "C"}',
            "headers": '{"performers": "On stage"}',
        }


def test_an_unverified_proposal_is_marked_before_its_defaults_are_used(
    tmp_path, capsys, monkeypatch
):
    """Spec section 4: shown marked unverified for hand-editing. The
    marker must come before the first question, because that question's
    default is the unverified value."""
    _fake_provider(monkeypatch, UnverifiedProvider())
    monkeypatch.chdir(tmp_path)
    seen = []

    def note(text=""):
        seen.append(text)

    code = run_wizard(
        VIVO,
        input_fn=lambda prompt: (seen.append(prompt), "")[1],
        print_fn=note,
        output=str(tmp_path / "out.yaml"),
        force=False,
        scene=None,
        one_shot=False,
    )
    markers = [i for i, s in enumerate(seen) if "proposal unverified" in s]
    assert len(markers) == 1
    assert "title" in seen[markers[0]]
    first_question = next(
        i for i, s in enumerate(seen) if s.startswith("Sheet")
    )
    assert markers[0] < first_question
    # Without a title the mapping is refused before anything is saved.
    assert code == 1


def test_provider_failure_prints_one_line_and_walks_manual_questions(
    tmp_path, capsys, monkeypatch
):
    _fake_provider(monkeypatch, DeadProvider())
    monkeypatch.chdir(tmp_path)
    answers = iter([
        "Rundown",       # sheet, typed by hand -- nothing to accept
        "4",             # header row
        "B", "C", "F",   # id, time, title by letter
        "On stage",      # performers by header text
        "CHUẨN BỊ",      # note
        "", "", "", "",  # sound, lighting, led left unmapped
    ])
    out = tmp_path / "vivo.yaml"
    code = run_wizard(
        VIVO,
        input_fn=lambda *a, **k: next(answers),
        print_fn=lambda *a, **k: print(*a),
        output=str(out),
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 0
    lines = [
        line
        for line in capsys.readouterr().out.splitlines()
        if "no model assist" in line
    ]
    assert len(lines) == 1
    assert out.exists()
    assert "segments:" in out.read_text(encoding="utf-8")


def test_unresolved_performers_are_offered_a_guess_after_the_render(
    tmp_path, capsys, monkeypatch
):
    """J2 wiring: after the show-context renders, each term the vocabulary
    could not read is offered a model-proposed kind. The answer here is
    'n', so this test must never touch any classifier.yaml."""
    from openpyxl import Workbook

    _fake_provider(monkeypatch, BothProvider())
    monkeypatch.chdir(tmp_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rundown"
    ws.append(["No", "Nội dung", "On stage"])
    ws.append([1, "Đón khách", "tốp múa"])
    xlsx = tmp_path / "run.xlsx"
    wb.save(xlsx)

    asked = []

    def answers(prompt):
        asked.append(prompt)
        return "n" if prompt.startswith("Record") else ""

    out = tmp_path / "out.yaml"
    code = run_wizard(
        str(xlsx),
        input_fn=answers,
        print_fn=lambda *a, **k: None,
        output=str(out),
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 0
    record = [p for p in asked if p.startswith("Record")]
    assert len(record) == 1
    assert "'tốp múa'" in record[0] and "speech.playback" in record[0]


def test_cli_one_shot_flag_reaches_the_wizard(tmp_path, capsys, monkeypatch):
    """Wiring check: the flag travels into run_wizard. The wizard body
    itself (Enter-to-accept, one map file) is proven directly above --
    driving it through main() would read stdin under capsys."""
    seen = {}

    def fake_wizard(xlsx, **kwargs):
        seen.update(kwargs, xlsx=xlsx)
        return 0

    monkeypatch.setattr(
        "wing_parser.showcontext.ingest.wizard.run_wizard", fake_wizard
    )
    code = main(["showcontext", "import", VIVO, "--one-shot"])
    assert code == 0
    assert seen["xlsx"] == VIVO
    assert seen["one_shot"] is True


def test_map_with_one_shot_is_refused(tmp_path, capsys):
    mapping = tmp_path / "map.yaml"
    mapping.write_text("header_row: 1\ncolumns:\n  title: A\n", encoding="utf-8")
    code = main([
        "showcontext", "import", VIVO,
        "--map", str(mapping), "--one-shot",
    ])
    assert code == 2
    err = capsys.readouterr().err
    assert "Traceback" not in err
    assert "--one-shot" in err


def test_no_assist_without_a_mapping_is_refused(tmp_path, capsys, monkeypatch):
    monkeypatch.delenv("WING_PROVIDER_CONFIG", raising=False)
    code = main(["showcontext", "import", VIVO, "--no-assist"])
    assert code != 0
    assert "--no-assist" in capsys.readouterr().err


def test_wizard_refuses_to_overwrite_output_without_force(tmp_path, monkeypatch):
    _fake_provider(monkeypatch, VivoProvider())
    monkeypatch.chdir(tmp_path)
    out = tmp_path / "vivo.yaml"
    out.write_text("keep me\n", encoding="utf-8")
    asked = []
    code = run_wizard(
        VIVO,
        input_fn=lambda prompt: (asked.append(prompt), "")[1],
        print_fn=lambda *a, **k: None,
        output=str(out),
        force=False,
        scene=None,
        one_shot=False,
    )
    assert code == 1
    assert out.read_text(encoding="utf-8") == "keep me\n"
    assert not asked  # refused before a single question was asked


# --- D-39: a key saved in the desktop Settings dialog must reach the CLI ---

SAVED_BY_SETTINGS = (
    "provider: openai-compat\n"
    "model: deepseek-chat\n"
    "base_url: https://api.deepseek.com\n"
    "api_key_env: DEEPSEEK_API_KEY\n"
    "api_key: sk-saved-from-settings\n"
)


def _recording_factory(monkeypatch, provider):
    """Swap make_provider for one that records every config handed to it."""
    built = []

    def record(config):
        built.append(config)
        return provider

    monkeypatch.setattr(
        "wing_parser.classifier.provider.make_provider", record
    )
    return built


def _no_key_environment(monkeypatch, tmp_path):
    """No pin, no ./provider.yaml, no provider env keys -- the operator's
    machine right after they pasted the key into Tools > Settings."""
    monkeypatch.delenv("WING_PROVIDER_CONFIG", raising=False)
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.chdir(tmp_path)
    assert not (tmp_path / "provider.yaml").exists()


def test_the_wizard_reads_the_key_settings_saved_in_the_knowledge_dir(
    tmp_path, monkeypatch
):
    """D-39: the copy `SettingsDialog.save()` writes is what the wizard's
    provider factory builds with, in a process that never saw the dialog's
    `$WING_PROVIDER_CONFIG` pin. The mapping-proposal site; the term-guess
    site is pinned by the test below."""
    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    (knowledge / "provider.yaml").write_text(
        SAVED_BY_SETTINGS, encoding="utf-8"
    )
    _no_key_environment(monkeypatch, tmp_path)
    built = _recording_factory(monkeypatch, VivoProvider())

    code = run_wizard(
        VIVO,
        input_fn=lambda *a, **k: "",
        print_fn=lambda *a, **k: None,
        output=str(tmp_path / "vivo.yaml"),
        force=False,
        scene=None,
        one_shot=True,
        knowledge_dir=knowledge,
    )
    assert code == 0
    assert [config.api_key for config in built] == ["sk-saved-from-settings"]


def test_without_a_knowledge_dir_the_wizard_still_resolves_the_old_way(
    tmp_path, monkeypatch
):
    """The default stays `load_config(None)`: a knowledge-dir copy that was
    never passed in must not be found by some ambient lookup."""
    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    (knowledge / "provider.yaml").write_text(
        SAVED_BY_SETTINGS, encoding="utf-8"
    )
    _no_key_environment(monkeypatch, tmp_path)
    built = _recording_factory(monkeypatch, VivoProvider())

    code = run_wizard(
        VIVO,
        input_fn=lambda *a, **k: "",
        print_fn=lambda *a, **k: None,
        output=str(tmp_path / "vivo.yaml"),
        force=False,
        scene=None,
        one_shot=True,
    )
    assert code == 0
    assert [config.api_key for config in built] == [""]


def test_the_term_guess_reads_the_same_saved_key(tmp_path, monkeypatch):
    """The wizard's second model site -- `guess.offer_terms`' provider
    factory -- takes the knowledge dir too, so the two cannot disagree
    inside one run. Modelled on the J2 wiring test above; the answer is
    'n', so no classifier.yaml is touched."""
    from openpyxl import Workbook

    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    (knowledge / "provider.yaml").write_text(
        SAVED_BY_SETTINGS, encoding="utf-8"
    )
    _no_key_environment(monkeypatch, tmp_path)
    built = _recording_factory(monkeypatch, BothProvider())

    wb = Workbook()
    ws = wb.active
    ws.title = "Rundown"
    ws.append(["No", "Nội dung", "On stage"])
    ws.append([1, "Đón khách", "tốp múa"])
    xlsx = tmp_path / "run.xlsx"
    wb.save(xlsx)

    asked = []

    def answers(prompt):
        asked.append(prompt)
        return "n" if prompt.startswith("Record") else ""

    code = run_wizard(
        str(xlsx),
        input_fn=answers,
        print_fn=lambda *a, **k: None,
        output=str(tmp_path / "out.yaml"),
        force=False,
        scene=None,
        one_shot=False,
        knowledge_dir=knowledge,
    )
    assert code == 0
    # Both sites ran: the mapping proposal, then the term guess.
    assert [p for p in asked if p.startswith("Record")]
    assert len(built) == 2
    assert {config.api_key for config in built} == {"sk-saved-from-settings"}


def test_cli_passes_the_knowledge_dir_to_the_wizard(tmp_path, monkeypatch):
    """Wiring check, modelled on the --one-shot one above: `showcontext
    import` resolves the knowledge directory and hands it to run_wizard."""
    seen = {}

    def fake_wizard(xlsx, **kwargs):
        seen.update(kwargs, xlsx=xlsx)
        return 0

    monkeypatch.setattr(
        "wing_parser.showcontext.ingest.wizard.run_wizard", fake_wizard
    )
    monkeypatch.setenv("WING_KNOWLEDGE_DIR", str(tmp_path))
    code = main(["showcontext", "import", VIVO, "--one-shot"])
    assert code == 0
    assert seen["knowledge_dir"] == Path(tmp_path)
