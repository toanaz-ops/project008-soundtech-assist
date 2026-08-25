"""Spreadsheet to cell text. This module interprets nothing.

It is the only module that knows openpyxl exists, which is what makes a
later CSV, PDF or photo reader a sibling of this file rather than a
change to the three modules downstream.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path

EXTRA_HINT = (
    "reading a spreadsheet needs openpyxl, which is not installed. "
    "Install it with:  pip install -e .[ingest]"
)


class MissingExtra(RuntimeError):
    """openpyxl is absent, so there is nothing to read the sheet with.

    This raises, and names the extra to install. It is the opposite of
    how `classifier/llm.py` treats a missing `anthropic`: `available()`
    there swallows the ImportError and returns False
    (`classifier/llm.py:60-69`), because the model is an optional
    fallback and the tool is designed to work without it. A spreadsheet
    reader is not optional to a spreadsheet reader, so an absent one is
    a stop, not a quieter path.
    """


@dataclass(frozen=True)
class RawRow:
    number: int
    cells: dict[str, str]


@dataclass(frozen=True)
class SheetRead:
    """`last_column` is the letter of the last column holding anything.

    `headers` drops a column whose header cell is blank, and an
    unlabelled notes or STT column is common on a real running order --
    exactly the case `headers:` cannot express and letters exist for
    (design spec section 4.3). Reporting the sheet's extent separately is
    what lets `mapping.resolve_columns` accept such a letter while still
    refusing one past the end of the sheet. Empty only for a sheet with
    nothing in it at all.
    """

    headers: dict[str, str]
    rows: tuple[RawRow, ...]
    blank_rows: int
    last_column: str = ""


def _text(value) -> str:
    """Every cell becomes stripped text.

    A time cell arrives as datetime.time, a numbered column as int. Both
    become text here because nothing downstream does arithmetic on a cell
    -- see the design spec section 5.1 on why times are not converted.
    """
    if value is None:
        return ""
    return str(value).strip()


def worksheet_for(book, sheet: str | int | None):
    if sheet is None:
        return book.worksheets[0]
    if isinstance(sheet, int):
        if not 1 <= sheet <= len(book.worksheets):
            raise ValueError(
                f"sheet index {sheet} is out of range; the workbook has "
                f"{len(book.worksheets)} sheet(s)."
            )
        return book.worksheets[sheet - 1]
    if sheet not in book.sheetnames:
        raise ValueError(
            f"no sheet named {sheet!r}; this workbook has: "
            + ", ".join(repr(name) for name in book.sheetnames)
        )
    return book[sheet]


# Backwards-compatible alias for internal callers.
_worksheet = worksheet_for


def read_sheet(path: str | Path, sheet: str | int | None, header_row: int) -> SheetRead:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:                # pragma: no cover - env-dependent
        raise MissingExtra(EXTRA_HINT) from exc

    if header_row < 1:
        raise ValueError(f"header_row must be 1 or more, not {header_row}.")

    # A file that is not a real .xlsx -- wrong extension, a renamed .docx,
    # plain text saved with an .xlsx name -- surfaces as zipfile.BadZipFile
    # (not a zip at all) or a KeyError (a zip, but missing the parts every
    # xlsx has), both from deep inside openpyxl/zipfile. Neither is a
    # ValueError the CLI already catches, so both are named here and
    # turned into the same shape: an error naming the file and what was
    # expected, the same convention `mapping.load_mapping` follows for a
    # broken YAML file.
    try:
        book = load_workbook(Path(path), data_only=True, read_only=True)
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError(f"{path}: not a readable .xlsx file ({exc}).") from exc
    try:
        worksheet = worksheet_for(book, sheet)
        # Materialised inside the try: a read-only workbook cannot be
        # iterated once it is closed.
        raw = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        book.close()

    grid = [
        {get_column_letter(index): _text(cell)
         for index, cell in enumerate(row, start=1)}
        for row in raw
    ]
    widest = max(
        (index
         for row in raw
         for index, cell in enumerate(row, start=1)
         if _text(cell)),
        default=0,
    )
    last_column = get_column_letter(widest) if widest else ""

    if header_row > len(grid):
        raise ValueError(
            f"header_row {header_row} is past the end of the sheet, which has "
            f"{len(grid)} row(s)."
        )

    headers = {letter: text for letter, text in grid[header_row - 1].items() if text}
    if not headers:
        raise ValueError(f"row {header_row} is empty, so it cannot be the header row.")

    rows: list[RawRow] = []
    blank = 0
    for offset, cells in enumerate(grid[header_row:], start=header_row + 1):
        if any(cells.values()):
            rows.append(RawRow(number=offset, cells=cells))
        else:
            blank += 1

    return SheetRead(
        headers=headers,
        rows=tuple(rows),
        blank_rows=blank,
        last_column=last_column,
    )
